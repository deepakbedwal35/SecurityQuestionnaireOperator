"""
File operations tools: find inbox files, parse questionnaire xlsx, and write filled xlsx.
"""
import glob
import os
from typing import Any, Dict, List, Optional
from pydantic import BaseModel, Field
import openpyxl

from agent.policy import PolicyEngine
from agent.state import Answer, Question, QuestionType, RiskLevel, RunState
from agent.tools.registry import Tool


# ---- Tool: find_inbox_files ----

class FindInboxFilesInput(BaseModel):
    inbox_dir: str = Field(default="data/inbox", description="Directory to search for incoming files")
    pattern: str = Field(default="*.xlsx", description="File pattern to match")


class FindInboxFilesOutput(BaseModel):
    files: List[str] = Field(default_factory=list)


class FindInboxFilesTool(Tool):
    name = "find_inbox_files"
    description = "Searches the inbox directory for incoming questionnaires (xlsx)."
    risk = RiskLevel.READ
    args_schema = FindInboxFilesInput
    output_schema = FindInboxFilesOutput

    def run(self, state: RunState, **kwargs: Any) -> Dict[str, Any]:
        inbox_dir = kwargs.get("inbox_dir", "data/inbox")
        pattern = kwargs.get("pattern", "*.xlsx")
        search_path = os.path.join(inbox_dir, pattern)
        matched = glob.glob(search_path)
        return {"files": sorted(matched)}


# ---- Tool: parse_questionnaire ----

class ParseQuestionnaireInput(BaseModel):
    file_path: str = Field(..., description="Path to the questionnaire xlsx file")


class ParseQuestionnaireOutput(BaseModel):
    question_count: int
    questions: List[Question]


class ParseQuestionnaireTool(Tool):
    name = "parse_questionnaire"
    description = "Parses an Excel security questionnaire and extracts questions into structured models with policy checks."
    risk = RiskLevel.READ
    args_schema = ParseQuestionnaireInput
    output_schema = ParseQuestionnaireOutput

    def run(self, state: RunState, **kwargs: Any) -> Dict[str, Any]:
        file_path = kwargs["file_path"]
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Questionnaire file not found: {file_path}")

        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active

        extracted_questions: List[Question] = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not any(row):
                continue

            # Standard columns: col 1 = ID, col 2 = section, col 3 = text
            qid = str(row[0] or f"Q-{row_idx:02d}").strip()
            text_cell = str(row[2] or "").strip()
            if not text_cell:
                continue

            # Sanitize untrusted cell content
            sanitized_text = PolicyEngine.sanitize_untrusted_text(text_cell)
            q_type, flags = PolicyEngine.classify_question(sanitized_text)

            question = Question(
                id=qid,
                text=sanitized_text,
                type=q_type,
                raw_row=row_idx,
                flags=flags,
            )
            extracted_questions.append(question)

            # Record injection flags on state if detected (deduped by source + sha1(snippet))
            for flag in flags:
                if "prompt_injection_detected" in flag:
                    state.add_injection_flag(
                        source=f"{file_path}:row_{row_idx}",
                        snippet=sanitized_text[:100],
                        action_taken="Flagged and neutralized",
                    )

        state.questions = extracted_questions
        return {
            "question_count": len(extracted_questions),
            "questions": extracted_questions,
        }


# ---- Tool: write_filled_xlsx ----

class WriteFilledXlsxInput(BaseModel):
    source_file: str = Field(..., description="Original questionnaire xlsx")
    output_file: str = Field(..., description="Destination xlsx with answers filled")
    answers: Optional[List[Dict[str, Any]]] = Field(
        default=None,
        description="Optional explicit answers list, otherwise reads from state.answers",
    )


class WriteFilledXlsxOutput(BaseModel):
    output_path: str
    filled_count: int


class WriteFilledXlsxTool(Tool):
    name = "write_filled_xlsx"
    description = "Writes answers and citations back into the questionnaire Excel spreadsheet."
    risk = RiskLevel.WRITE_LOW
    args_schema = WriteFilledXlsxInput
    output_schema = WriteFilledXlsxOutput

    def run(self, state: RunState, **kwargs: Any) -> Dict[str, Any]:
        source_file = kwargs["source_file"]
        output_file = kwargs["output_file"]
        
        # Determine answers from argument or state
        answers_data = kwargs.get("answers")
        if answers_data is not None:
            answer_map = {a["question_id"]: a for a in answers_data}
        else:
            answer_map = {a.question_id: a.model_dump() for a in state.answers}

        wb = openpyxl.load_workbook(source_file)
        ws = wb.active

        filled_count = 0
        for row_idx in range(2, ws.max_row + 1):
            qid_val = ws.cell(row=row_idx, column=1).value
            if not qid_val:
                continue
            qid = str(qid_val).strip()
            if qid in answer_map:
                ans_record = answer_map[qid]
                ans_text = ans_record.get("text", "")
                decision = ans_record.get("decision", "answered")

                if decision == "escalated":
                    cell_text = f"[ESCALATED FOR APPROVAL] {ans_text}"
                elif decision == "flagged":
                    cell_text = f"[SECURITY FLAGGED] {ans_text}"
                else:
                    cell_text = ans_text

                citations_list = ans_record.get("citations", [])
                citations_text = "; ".join(
                    f"{c.get('doc')} (quote: \"{c.get('quote')}\")" for c in citations_list if isinstance(c, dict)
                )

                ws.cell(row=row_idx, column=4, value=cell_text)
                ws.cell(row=row_idx, column=5, value=citations_text)
                filled_count += 1

        os.makedirs(os.path.dirname(os.path.abspath(output_file)), exist_ok=True)
        wb.save(output_file)
        state.artifacts["filled_xlsx"] = output_file

        return {
            "output_path": output_file,
            "filled_count": filled_count,
        }

