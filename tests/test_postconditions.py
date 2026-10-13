"""
Unit tests for each postcondition kind in Verifier (passing and failing cases).
"""
import os
import tempfile
import pytest
import openpyxl

from agent.state import (
    Answer,
    AnswerDecision,
    Citation,
    PostconditionSpec,
    Question,
    QuestionType,
    RunState,
)
from agent.verifier import Verifier


def test_verify_file_exists():
    state = RunState(run_id="test", request_text="test")

    # Failing case: non-existent file
    spec_fail = PostconditionSpec(kind="file_exists", params={"path": "/nonexistent/path/file.xlsx"})
    res_fail = Verifier.verify(spec_fail, state)
    assert res_fail.passed is False

    # Passing case: existing non-empty file
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as tf:
        tf.write(b"content")
        path = tf.name

    try:
        spec_pass = PostconditionSpec(kind="file_exists", params={"path": path})
        res_pass = Verifier.verify(spec_pass, state)
        assert res_pass.passed is True
    finally:
        if os.path.exists(path):
            os.remove(path)


def test_verify_questions_parsed():
    state = RunState(run_id="test", request_text="test")
    spec = PostconditionSpec(kind="questions_parsed", params={"min_count": 2})

    # Failing case: empty questions
    res_fail = Verifier.verify(spec, state)
    assert res_fail.passed is False

    # Passing case
    state.questions = [
        Question(id="Q-01", text="Question 1", type=QuestionType.FACTUAL, raw_row=2),
        Question(id="Q-02", text="Question 2", type=QuestionType.FACTUAL, raw_row=3),
    ]
    res_pass = Verifier.verify(spec, state)
    assert res_pass.passed is True


def test_verify_all_questions_have_decision():
    state = RunState(run_id="test", request_text="test")
    state.questions = [
        Question(id="Q-01", text="Question 1", type=QuestionType.FACTUAL, raw_row=2),
        Question(id="Q-02", text="Question 2", type=QuestionType.FACTUAL, raw_row=3),
    ]
    spec = PostconditionSpec(kind="all_questions_have_decision", params={})

    # Failing case: 1 answer missing
    state.answers = [
        Answer(question_id="Q-01", text="Answer 1", decision=AnswerDecision.ANSWERED),
    ]
    res_fail = Verifier.verify(spec, state)
    assert res_fail.passed is False

    # Passing case: both have decisions
    state.answers.append(
        Answer(question_id="Q-02", text="Escalated", decision=AnswerDecision.ESCALATED)
    )
    res_pass = Verifier.verify(spec, state)
    assert res_pass.passed is True


def test_verify_no_commitment_question_answered():
    state = RunState(run_id="test", request_text="test")
    state.questions = [
        Question(
            id="Q-SLA",
            text="Do you guarantee 99.99% monthly service uptime with financial credits?",
            type=QuestionType.COMMITMENT,
            raw_row=2,
        )
    ]
    spec = PostconditionSpec(kind="no_commitment_question_answered", params={})

    # Failing case: commitment question answered auto-compliantly
    state.answers = [
        Answer(question_id="Q-SLA", text="Yes guaranteed.", decision=AnswerDecision.ANSWERED)
    ]
    res_fail = Verifier.verify(spec, state)
    assert res_fail.passed is False
    assert "incorrectly answered" in res_fail.detail.lower()

    # Passing case: escalated
    state.answers[0].decision = AnswerDecision.ESCALATED
    res_pass = Verifier.verify(spec, state)
    assert res_pass.passed is True


def test_verify_every_answered_has_citation():
    state = RunState(run_id="test", request_text="test")
    spec = PostconditionSpec(kind="every_answered_has_citation", params={})

    # Failing case: answered without citation
    state.answers = [
        Answer(question_id="Q-01", text="MFA is on.", decision=AnswerDecision.ANSWERED, citations=[])
    ]
    res_fail = Verifier.verify(spec, state)
    assert res_fail.passed is False

    # Passing case: citation present
    state.answers[0].citations = [Citation(doc="Policy.md", quote="MFA is mandatory")]
    res_pass = Verifier.verify(spec, state)
    assert res_pass.passed is True


def test_verify_output_xlsx_complete():
    state = RunState(run_id="test", request_text="test")

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
        xlsx_path = tf.name

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="ID")
    ws.cell(row=1, column=4, value="Answer")
    ws.cell(row=2, column=1, value="Q-01")
    ws.cell(row=2, column=4, value="Filled answer")
    ws.cell(row=3, column=1, value="Q-02")
    ws.cell(row=3, column=4, value="")  # Empty answer row
    wb.save(xlsx_path)

    try:
        spec = PostconditionSpec(kind="output_xlsx_complete", params={"output_path": xlsx_path})
        # Failing case: row 3 is empty
        res_fail = Verifier.verify(spec, state)
        assert res_fail.passed is False

        # Passing case: fill row 3
        ws.cell(row=3, column=4, value="Answer 2")
        wb.save(xlsx_path)
        res_pass = Verifier.verify(spec, state)
        assert res_pass.passed is True
    finally:
        if os.path.exists(xlsx_path):
            os.remove(xlsx_path)
